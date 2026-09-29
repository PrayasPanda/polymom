"""MoM exports for a code-mixed (English, Hindi, Odia) meeting."""

import io
import json
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.core.config import Settings
from app.schemas.meeting import AudioMetadata, MeetingRead, MeetingStatus
from app.schemas.result import MeetingResult, SpeakerInfo
from app.services.export import render_export
from app.services.export.document import build_document
from app.services.llm.mock_client import MockLLMClient
from app.services.summarization.summarizer import Summarizer
from scripts.eval_summary import load_transcript

FIXTURE = Path(__file__).parents[1] / "fixtures" / "meetings" / "mixed_standup.json"
SECTIONS = [
    "Executive summary",
    "Decisions",
    "Action items",
    "Open questions",
    "Appendix: transcript",
]


@pytest.fixture
async def result() -> MeetingResult:
    fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
    transcript = load_transcript(fixture)
    settings = Settings(_env_file=None, llm_provider="mock")
    summary = await Summarizer(MockLLMClient(settings, "mock-llm"), settings).summarize(transcript)
    now = datetime(2026, 9, 28, 10, 0, tzinfo=UTC)
    meeting = MeetingRead(
        meeting_id=uuid.uuid4(),
        title="Daily standup",
        original_filename="standup.wav",
        mime_type="audio/x-wav",
        size_bytes=1,
        duration_seconds=36.0,
        audio_metadata=AudioMetadata(),
        languages_hint=[],
        expected_speakers=None,
        detected_languages=["en", "hi", "or"],
        status=MeetingStatus.COMPLETED,
        error=None,
        created_at=now,
        updated_at=now,
    )
    return MeetingResult(
        meeting=meeting,
        processing=None,
        speakers=[
            SpeakerInfo(label="Person 1", display_name="Ravi"),
            SpeakerInfo(label="Person 2", display_name="सुनीता"),
            SpeakerInfo(label="Person 3", display_name="ପ୍ରିୟା"),
        ],
        transcript=transcript,
        summary=summary,
        verification_report=summary.verification_report,
        included=["transcript", "summary"],
    )


def squash(text: str) -> str:
    """pypdf inserts spaces/control chars around shaped Indic clusters; compare without them."""
    return re.sub(r"[\s\x00-\x1f]+", "", text)


def test_document_view_model(result: MeetingResult) -> None:
    doc = build_document(result)
    assert doc.title == result.summary.title  # type: ignore[union-attr]
    assert doc.duration == "00:00:36"
    assert doc.languages == ["English", "Hindi", "Odia"]
    assert [p.name for p in doc.participants] == ["Ravi", "सुनीता", "ପ୍ରିୟା"]
    assert doc.decisions
    assert doc.action_items
    assert {row[2] for row in doc.action_items} <= {"Ravi", "सुनीता", "ପ୍ରିୟା", "Unassigned"}
    assert doc.transcript[1].speaker == "सुनीता"


async def test_pdf_has_sections_and_extractable_indic_text(result: MeetingResult) -> None:
    pdf = render_export("pdf", result)

    reader = PdfReader(io.BytesIO(pdf))
    text = "\n".join(page.extract_text() for page in reader.pages)
    fonts = {
        str(font.get("/BaseFont"))
        for page in reader.pages
        for font in (page["/Resources"].get("/Font") or {}).values()  # type: ignore[union-attr]
        for font in [font.get_object()]
    }
    assert len(reader.pages) >= 2  # transcript appendix starts on a new page
    for section in SECTIONS:
        assert section in text
    flat = squash(text)
    for word in ("timeout", "सुनीता", "देख", "ପ୍ରିୟା", "କରିବି"):
        assert squash(word) in flat, word
    assert any("NotoSansDevanagari" in f for f in fonts)
    assert any("NotoSansOriya" in f for f in fonts)


async def test_docx_has_sections_tables_and_scripts(result: MeetingResult) -> None:
    document = Document(io.BytesIO(render_export("docx", result)))

    paragraphs = [p.text for p in document.paragraphs]
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert all(section in headings for section in SECTIONS)
    assert document.tables
    assert document.tables[0].rows[0].cells[1].text in {"Decision", "Task"}
    body = "\n".join(paragraphs)
    assert "सुनीता" in body
    assert "ମୁଁ agle hafte load testing କରିବି।" in body


async def test_markdown_and_json(result: MeetingResult) -> None:
    md = render_export("md", result).decode()
    assert md.startswith("# ")
    for section in SECTIONS:
        assert f"## {section}" in md
    assert "**Participants:** Ravi; सुनीता; ପ୍ରିୟା" in md
    assert json.loads(render_export("json", result))["schema_version"] == "1.0.0"


async def test_exports_without_summary(result: MeetingResult) -> None:
    bare = result.model_copy(update={"summary": None, "transcript": None})
    md = render_export("md", bare).decode()
    assert "_Summary not available._" in md
    assert "_Transcript not available._" in md
    assert render_export("pdf", bare).startswith(b"%PDF")
    assert render_export("docx", bare).startswith(b"PK")
